"""Create an enlarged nine-page Word editing draft from the research dossier."""
from pathlib import Path
import json
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "manuscript"
OUT.mkdir(exist_ok=True)
DOCX = OUT / "Iron_Web_Research_Paper_9_Page_Editing_Draft.docx"

ABSTRACT = """Persistent memory can help language model agents maintain continuity, but retrieving a relevant passage does not establish whether its claims remain valid, supported, or applicable to the current task. This paper presents Iron Web, an architecture that combines retained source text, vector retrieval, a temporal claim graph, and bounded context assembly. Source records distinguish user statements from agent assertions, while explicit lifecycle operations represent confirmation, correction, supersession, and retraction. Retrieval reconciles evidence with temporal validity and unresolved conflicts before constructing context for a consuming model. The architecture also associates context and dependent work with the claims on which they rely, allowing changed evidence to trigger revalidation. We analyze Mem0, Zep, and a critical study of MemPalace to position these mechanisms within existing memory research. Published benchmark results are separated by dataset, reader model, and metric, with particular attention to the distinction between retrieval recall and answer accuracy. We propose a controlled evaluation using LoCoMo, LongMemEval, lifecycle scenarios, and component ablations. Existing project reports provide limited engineering validation, but public benchmark experiments for Iron Web remain outstanding. The present findings establish an implementation grounded research framework and identify measurable reliability questions; they do not establish a benchmark accuracy improvement or state of the art performance."""
ABSTRACT = ABSTRACT.replace("can help", "helps").replace("retained source text", "source text").replace("a controlled evaluation", "evaluation")
assert len(ABSTRACT.split()) == 200


def make_diagram():
    im = Image.new("RGB", (2100, 720), "white")
    draw = ImageDraw.Draw(im)
    font = ImageFont.truetype(r"C:\Windows\Fonts\times.ttf", 49)
    boxes = [(30,55,650,285),(740,55,1360,285),(1450,55,2070,285),
             (30,425,650,655),(740,425,1360,655),(1450,425,2070,655)]
    labels = ["Messages and\ntimestamps", "Source and claim\nvalidation", "Graph and vector\nstorage",
              "Grounded answer\nwith source links", "Bounded context\nand evidence checks", "Scoped hybrid\nretrieval"]
    for box, label in zip(boxes, labels):
        draw.rectangle(box, outline="black", width=4)
        draw.multiline_text(((box[0]+box[2])/2,(box[1]+box[3])/2),label,
                            font=font,fill="black",anchor="mm",align="center",spacing=14)
    def arrow(a,b):
        draw.line([a,b],fill="black",width=6)
        x,y=b; dx=b[0]-a[0];dy=b[1]-a[1]
        if dx>0: pts=[(x,y),(x-25,y-15),(x-25,y+15)]
        elif dx<0: pts=[(x,y),(x+25,y-15),(x+25,y+15)]
        else: pts=[(x,y),(x-15,y-25),(x+15,y-25)]
        draw.polygon(pts,fill="black")
    arrow((650,170),(740,170)); arrow((1360,170),(1450,170))
    arrow((1760,285),(1760,425)); arrow((1450,540),(1360,540));arrow((740,540),(650,540))
    path=OUT/"iron_web_architecture.png"
    im.save(path,dpi=(300,300))
    return path


doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65)
sec.left_margin=Inches(.75);sec.right_margin=Inches(.75)
sec.header_distance=Inches(.25);sec.footer_distance=Inches(.25)

for name in ["Normal","Title","Subtitle","Heading 1","Heading 2","Caption"]:
    s=doc.styles[name];s.font.name="Times New Roman";s.font.color.rgb=RGBColor(0,0,0)
    fonts=s.element.get_or_add_rPr().rFonts
    for attr in ["asciiTheme","hAnsiTheme","eastAsiaTheme","cstheme"]:
        fonts.attrib.pop(qn("w:"+attr),None)
    for attr in ["ascii","hAnsi","eastAsia","cs"]:
        fonts.set(qn("w:"+attr),"Times New Roman")
normal=doc.styles["Normal"]
normal.font.size=Pt(15)
normal.paragraph_format.line_spacing=1.08
normal.paragraph_format.space_after=Pt(8)
normal.paragraph_format.widow_control=True
normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
for name,size in [("Heading 1",18),("Heading 2",16.5)]:
    s=doc.styles[name];s.font.size=Pt(size);s.font.bold=True
    s.paragraph_format.space_before=Pt(5);s.paragraph_format.space_after=Pt(9)
    s.paragraph_format.keep_with_next=True
doc.styles["Title"].font.size=Pt(36)
doc.styles["Title"].font.bold=False
doc.styles["Title"].paragraph_format.space_after=Pt(12)
doc.styles["Title"].paragraph_format.line_spacing=1.0
doc.styles["Caption"].font.size=Pt(12)
doc.styles["Caption"].font.italic=False
doc.styles["Caption"].paragraph_format.space_after=Pt(7)
doc.styles["Caption"].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER

# Remove any inherited paragraph rule from the title.
for sname in ["Title","Subtitle"]:
    ppr=doc.styles[sname].element.find(qn("w:pPr"))
    if ppr is not None:
        for border in list(ppr.findall(qn("w:pBdr"))):ppr.remove(border)

hp=sec.header.paragraphs[0]
hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
hr=hp.add_run("IRON WEB  |  EXPANDED EDITING DRAFT");hr.font.name="Times New Roman";hr.font.size=Pt(9)
fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=fp.add_run();r.font.size=Pt(10)
fld=OxmlElement("w:fldSimple");fld.set(qn("w:instr"),"PAGE");r._r.addnext(fld)


def p(text, *, size=None, bold=False, italic=False, align=None, after=None):
    para=doc.add_paragraph()
    run=para.add_run(text);run.bold=bold;run.italic=italic
    if size:run.font.size=Pt(size)
    if align is not None:para.alignment=align
    if after is not None:para.paragraph_format.space_after=Pt(after)
    return para


NEXT_PAGE = False
def h(text,level=1):
    global NEXT_PAGE
    para=doc.add_paragraph(text,f"Heading {level}")
    if NEXT_PAGE:
        para.paragraph_format.page_break_before=True
        NEXT_PAGE=False
    return para
def newpage():
    global NEXT_PAGE
    NEXT_PAGE=True


def table(caption,headers,rows,widths):
    cap=doc.add_paragraph(caption,"Caption");cap.paragraph_format.keep_with_next=True
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for i,w in enumerate(widths):t.columns[i].width=Inches(w)
    for cell,text,w in zip(t.rows[0].cells,headers,widths):cell.text=text;cell.width=Inches(w)
    for row in rows:
        cells=t.add_row().cells
        for cell,text,w in zip(cells,row,widths):cell.text=str(text);cell.width=Inches(w)
    for ri,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr();cant=OxmlElement("w:cantSplit");pr.append(cant)
        if ri==0:repeat=OxmlElement("w:tblHeader");pr.append(repeat)
        for ci,cell in enumerate(row.cells):
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            margins=OxmlElement("w:tcMar")
            for side,amount in [("top",65),("bottom",65),("left",85),("right",85)]:
                e=OxmlElement("w:"+side);e.set(qn("w:w"),str(amount));e.set(qn("w:type"),"dxa");margins.append(e)
            tcpr.append(margins)
            if ri==0:
                shade=OxmlElement("w:shd");shade.set(qn("w:fill"),"EDEDED");tcpr.append(shade)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(0);para.paragraph_format.space_before=Pt(0)
                para.paragraph_format.line_spacing=1.0
                para.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:run.font.size=Pt(12);run.bold=(ri==0);run.font.color.rgb=RGBColor(0,0,0)
    borders=OxmlElement("w:tblBorders")
    for side in ["top","bottom","insideH"]:
        e=OxmlElement("w:"+side);e.set(qn("w:val"),"single");e.set(qn("w:sz"),"4");e.set(qn("w:color"),"777777");borders.append(e)
    t._tbl.tblPr.append(borders)
    spacer=p("",size=3,after=3);spacer.paragraph_format.line_spacing=Pt(3)
    return t


# PAGE 1
title=doc.add_paragraph("Iron Web\nTemporal Memory\nwith Source Evidence","Title")
title.alignment=WD_ALIGN_PARAGRAPH.CENTER
p("Author Name\nDepartment and Institution\nCity Country  |  Email Address",size=12,align=WD_ALIGN_PARAGRAPH.CENTER,after=12)
h("Abstract")
ap=p(ABSTRACT);ap.paragraph_format.line_spacing=1.0
p("Keywords: agent memory, temporal knowledge graphs, evidence provenance",size=15,after=5)
p("Draft status: architecture and literature analysis; public benchmark evaluation pending.",size=11,italic=True,after=0)

# PAGE 2
newpage();h("1. Introduction")
p("Language model agents need access to information that persists beyond a single interaction. Retrieval-augmented generation supplies external passages to a model [1], and memory hierarchies extend the amount of information that can be used despite a bounded context window [2]. However, persistence alone does not determine whether a remembered statement is still applicable. A user may correct an earlier assertion, change a preference, or describe different requirements for separate projects. Returning a semantically similar record can therefore produce a confident but outdated answer.")
p("Existing memory architectures address parts of this problem. Mem0 extracts and updates compact memories [3], while Zep constructs a temporal knowledge graph with retained episodes and evolving relationships [4]. A critical analysis of MemPalace argues that strong retrieval can also arise from retained text and ordinary vector search [5]. Together, these approaches motivate a design that preserves original evidence while explicitly representing changes to derived claims.")
p("Iron Web combines source records, structured claims, vector retrieval, graph traversal, and context assembly. A source statement and an accepted fact are separate objects. Confirmation, correction, supersession, and retraction have explicit meanings. When current evidence is disputed or insufficient, the system can expose that state rather than selecting a value solely because it ranks highly. Context and dependent work can also be associated with the claims they use, allowing later changes to trigger revalidation.")
p("This paper develops that architecture and a protocol for evaluating its reliability. LoCoMo [6] and LongMemEval [7] provide established conversational tests, but public benchmark accuracy must be supplemented with scenarios involving corrections, source attribution, and stale context. The intended contribution is an evaluated connection between claim lifecycle decisions and downstream memory use. Graph storage, temporal records, and tiered context are established techniques and are not claimed as independent inventions.")
p("The present draft reports a literature-based comparison and an inspection of the existing implementation. It does not report a newly executed public benchmark. The proposed experiments will determine whether the additional lifecycle mechanisms improve answer quality and operational reliability at a comparable context budget.")

# PAGE 3
newpage();h("2. Related Work")
p("Memory systems differ in what they retain and when they perform interpretation. MemoryBank applies updating and forgetting mechanisms to support persistent interaction [8]. ReadAgent stores compact gist memories while retaining the ability to consult original passages [9], and RAPTOR retrieves across a hierarchy of recursively summarized text [10]. These approaches show that compression and access to detailed evidence can coexist. The relevant design question is which information should be transformed, and how the original evidence remains recoverable.")
p("Graph approaches organize relationships across records. GraphRAG constructs entity graphs and community summaries for broad questions over document collections [11]. LightRAG combines graph structures with vector representations and incremental updates [12]. AriGraph connects episodic and semantic memory for agents interacting with environments [13]. Consequently, combining graph and vector retrieval is substantial prior art; Iron Web must demonstrate the value of its particular lifecycle and context behavior.")
p("A-MEM dynamically links structured notes and updates representations of earlier memories [14]. Reflexion uses persistent verbal feedback to improve later attempts [15]. Hindsight distinguishes several forms of retained knowledge and supports retain, recall, and reflect operations [16]. These systems are relevant to the separation of evidence, agent experience, and inference. An agent-generated reflection should not automatically receive the authority of a user-confirmed statement.")
p("The supplied Mem0 study reports LoCoMo quality and efficiency, while the Zep study evaluates Deep Memory Retrieval and LongMemEval. Their percentages are conditional on different readers, datasets, and protocols. The MemPalace analysis is useful for motivating raw-text baselines, but its cross-system table should not be reused as a uniform accuracy ranking. Retrieval recall and judged answer correctness measure different stages of the pipeline.")
p("Later project reports describe observational memory [17] and an updated temporal Mem0 algorithm [18]. LongMemEval-V2 extends the evaluation setting to environment experience and workflows [19]. These sources broaden the comparison, but their newer algorithms and tasks must remain distinct from the supplied 2025 experiments. This draft uses the related literature to define testable comparisons rather than infer a current state-of-the-art ranking.")

# PAGE 4
newpage();h("3. Proposed Methodology")
h("3.1 Source Records and Temporal Claims",2)
p("Iron Web stores original memory text alongside structured claims. Each source records its origin, actor, recording time, confirmation state, and evidence group. Each claim identifies a subject, predicate, value, scope, and validity interval. The distinction allows a statement to remain available as historical evidence even when it should no longer support a current answer. Source authority is supplied through a trusted local API; the content of a document does not grant itself authority.")
p("Lifecycle operations include assertion, confirmation, correction, supersession, and retraction. A correction records that an earlier assertion was mistaken, whereas supersession represents a subsequent change. Queries about a past state and queries about what was believed at that time may therefore require different interpretations. Conflicting eligible values can remain disputed. Repeated agent assertions share an evidence group by default so that repetition is not counted as independent corroboration.")
h("3.2 Storage and Scoped Hybrid Retrieval",2)
p("The durable engine coordinates a FalkorDB graph with embedded Qdrant vectors. A SQLite receipt journal supports processing state and replay; it is not treated as the complete factual store. Writes use coordinated operations and compensation rather than one crash-atomic transaction across both databases. This implementation boundary must be retained in claims about durability and recovery.")
p("Retrieval gathers semantic candidates and bounded graph neighbors, applies temporal and scope constraints, reconciles claim states, and optionally reranks the remaining evidence. Ranking uses semantic relevance with graph, recency, importance, entity, historical, contradiction, and reliability signals. The fuzzy reliability component is a heuristic, not a calibrated probability of truth. Hard eligibility rules remain separate from ranking preferences.")
p("The inspected defaults use 384-dimensional BGE embeddings, up to 100 vector candidates, two graph hops, and a 50-node graph bound. Reranking is disabled by default. These values describe the current configuration, not experimentally optimized settings. Reproducing a run requires saving the effective configuration after service and environment overrides.")

# PAGE 5
newpage();h("3.3 Bounded Context and Evaluation Design",2)
p("The context runtime organizes a bounded set of essential fields and selected context cards over durable evidence. It uses the consuming model's tokenizer to enforce its budget. Retrieved source content is represented as data rather than application instructions. Context validity depends on the claims it includes; changed or expired evidence can require refreshing context and checking dependent work before completion.")
fig=doc.add_paragraph();fig.alignment=WD_ALIGN_PARAGRAPH.CENTER
fig.add_run().add_picture(str(make_diagram()),width=Inches(7))
fig.paragraph_format.space_after=Pt(3);fig.paragraph_format.keep_with_next=True
doc.add_paragraph("Fig. 1. Iron Web ingestion and retrieval flow. Corrections can trigger context and dependent-work revalidation. The receipt journal supports replay alongside the storage operations.","Caption")
p("Evaluation will use frozen LoCoMo and LongMemEval datasets, a fixed answer model, and a documented judge. Baselines will include no memory, full history where it fits, lexical retrieval, dense raw-text retrieval, structured claims, and the complete hybrid system. Each comparison will preserve the reader's token budget. Questions, answers, and labelled evidence must remain unavailable during ingestion and tuning.")
p("Separate ablations will disable graph expansion, temporal filtering, lifecycle handling, source grouping, and context caching. Targeted timelines will test correction, historical validity, unresolved conflict, retraction, cross-scope retrieval, and stale-work rejection. Exact state correctness and source attribution will supplement answer accuracy. Hand-written answer-derived claims will be restricted to explicitly labelled oracle experiments, outside the main end-to-end comparison.")
p("The current context schema includes travel-specific task fields. A general conversational adapter remains necessary before public benchmark evaluation. Dataset revision, extraction policy, model identifiers, prompts, failures, and index readiness will be retained with each run.")

# PAGE 6
newpage();h("4. Results & Discussion")
p("The available quantitative evidence consists of published results and earlier local validation reports. Table I reproduces selected LoCoMo results from the Mem0 study [3, Table 2]. These are author-reported measurements under that study's protocol; they are not experiments executed for Iron Web. The proposed method is included with an explicit unmeasured status.")
table("TABLE I\nReported LoCoMo results from the Mem0 study",["Method","Judge score (%)","Total p95 (s)"],[
    ("Full context","72.90 ± 0.19","17.117"),("Zep","65.99 ± 0.16","2.926"),
    ("Mem0","66.88 ± 0.15","1.440"),("Mem0g","68.44 ± 0.17","2.590"),
    ("Iron Web proposed method","Not measured","Not measured")],[3.25,2.05,1.7])
p("Within this experiment, Mem0 improves on the ChatGPT-memory baseline while full context has the highest overall judge score. Mem0g gains 1.56 percentage points over Mem0, with additional response latency. Graph augmentation also reduces some category scores, so its benefit cannot be assumed for every reasoning type.")
table("TABLE II\nReported LongMemEval S results from the Zep study",["Method and reader","QA score (%)","Latency (s)"],[
    ("Full context with GPT-4o-mini","55.4","31.3"),("Zep with GPT-4o-mini","63.8","3.20"),
    ("Full context with GPT-4o","60.2","28.9"),("Zep with GPT-4o","71.2","2.58"),
    ("Iron Web proposed method","Not measured","Not measured")],[3.8,1.65,1.55])
p("Table II preserves the reader-specific comparisons in Zep [4, Table 2]. The source labels its timing column as latency, not p95. These figures therefore cannot be directly combined with Table I's p95 measurements, and the two tables describe different datasets.")

# PAGE 7
newpage();h("4. Results & Discussion",1)
p("MemPalace's reported 96.6% recall_any@5 [5] indicates that at least one labelled evidence session appeared among five retrieved results. It does not establish that all required evidence was returned or that a reader answered correctly. An end-to-end accuracy claim requires answer generation and an appropriate evaluation procedure. Iron Web should therefore report retrieval hit rate, evidence coverage, and answer correctness separately.")
p("The existing project record reports 204 passing automated checks and 21 passing live space checks [20]. Earlier Nemotron validation records report 18 successful scenario checks across two models. These are small engineering checks covering selected storage, correction, conflict, restart, supersession, and replay behavior. They are historical reports, not a newly executed experiment, and their counts must not be added together as independent benchmark questions.")
table("TABLE III\nIron Web evidence status and planned measurements",["Evaluation target","Present evidence","Required measurement"],[
    ("Conversation QA","Not measured","LoCoMo and LongMemEval accuracy"),
    ("Lifecycle behavior","Selected local scenarios","Current and historical state correctness"),
    ("Context freshness","Implemented checks","Stale-context and stale-work acceptance"),
    ("Efficiency","Small historical observations","Matched-budget p50 and p95 cost and latency")],[2.1,2.2,2.7])
p("A fair evaluation must also control ingestion quality. Iron Web's default extraction policy is none. Supplying hand-curated claims to Iron Web while baselines process raw dialogue would privilege the proposed method. An automated extraction policy must be frozen for an end-to-end study, and raw-only variants should remain part of the comparison.")
p("Confidence intervals should account for questions sharing a conversation. LoCoMo's ten conversations limit the precision of cluster-based estimates. Judge disagreement, timeouts, indexing delays, and unsuccessful runs should remain visible. In addition to accuracy, experiments should measure ingestion cost, retrieval cost, context size, memory growth, and recovery behavior. Until those measurements exist, the evidence supports architectural feasibility and a testable research design, rather than a numerical performance advantage.")

# PAGE 8
newpage();h("5. Conclusion")
p("Iron Web provides an architecture for retaining source evidence while explicitly managing the lifecycle of derived claims. Its design connects temporal validity, source distinctions, scoped retrieval, bounded context, and dependent-work checks. These mechanisms address situations in which a relevant memory may still be outdated, disputed, or unsuitable for the current task.")
p("The literature comparison shows why memory research must separate retrieval recall from answer accuracy and preserve the dataset, model, and timing protocol behind every percentage. Existing local reports provide limited engineering evidence for Iron Web, while its public benchmark performance remains unmeasured. The next step is a frozen, reproducible evaluation with raw-text baselines, matched context budgets, lifecycle scenarios, and component ablations. The resulting evidence will determine whether the proposed integration improves reliability and whether its additional complexity is justified.")
h("References")

refs=[
    ('P. Lewis et al., “Retrieval-augmented generation for knowledge-intensive NLP tasks,” Advances in Neural Information Processing Systems, vol. 33, pp. 9459–9474, 2020.','https://arxiv.org/abs/2005.11401'),
    ('C. Packer et al., “MemGPT: Towards LLMs as operating systems,” arXiv preprint arXiv:2310.08560v2, 2024.','https://arxiv.org/abs/2310.08560v2'),
    ('P. Chhikara, D. Khant, S. Aryan, T. Singh, and D. Yadav, “Mem0: Building production-ready AI agents with scalable long-term memory,” arXiv:2504.19413v1, 2025.','https://arxiv.org/abs/2504.19413v1'),
    ('P. Rasmussen, P. Paliychuk, T. Beauvais, J. Ryan, and D. Chalef, “Zep: A temporal knowledge graph architecture for agent memory,” arXiv:2501.13956v1, 2025.','https://arxiv.org/abs/2501.13956v1'),
    ('R. Dey and P. Viradecha, “Spatial metaphors for LLM memory: A critical analysis of the MemPalace architecture,” arXiv:2604.21284v1, 2026.','https://arxiv.org/abs/2604.21284v1'),
    ('A. Maharana et al., “Evaluating very long-term conversational memory of LLM agents,” Proc. 62nd Annual Meeting of the Association for Computational Linguistics, pp. 13851–13870, 2024.','https://aclanthology.org/2024.acl-long.747/'),
    ('D. Wu et al., “LongMemEval: Benchmarking chat assistants on long-term interactive memory,” arXiv:2410.10813v2, 2025.','https://arxiv.org/abs/2410.10813v2'),
    ('W. Zhong, L. Guo, Q. Gao, H. Ye, and Y. Wang, “MemoryBank: Enhancing large language models with long-term memory,” arXiv:2305.10250v3, 2023.','https://arxiv.org/abs/2305.10250v3'),
    ('K.-H. Lee, X. Chen, H. Furuta, J. Canny, and I. Fischer, “A human-inspired reading agent with gist memory of very long contexts,” arXiv:2402.09727v3, 2024.','https://arxiv.org/abs/2402.09727v3'),
    ('P. Sarthi et al., “RAPTOR: Recursive abstractive processing for tree-organized retrieval,” arXiv:2401.18059v1, 2024.','https://arxiv.org/abs/2401.18059v1'),
    ('D. Edge et al., “From local to global: A graph RAG approach to query-focused summarization,” arXiv:2404.16130v2, 2025.','https://arxiv.org/abs/2404.16130v2'),
    ('Z. Guo, L. Xia, Y. Yu, T. Ao, and C. Huang, “LightRAG: Simple and fast retrieval-augmented generation,” arXiv:2410.05779v3, 2025.','https://arxiv.org/abs/2410.05779v3'),
    ('P. Anokhin et al., “AriGraph: Learning knowledge graph world models with episodic memory for LLM agents,” arXiv:2407.04363v3, 2025.','https://arxiv.org/abs/2407.04363v3'),
    ('W. Xu et al., “A-MEM: Agentic memory for LLM agents,” arXiv:2502.12110v11, 2025.','https://arxiv.org/abs/2502.12110v11'),
    ('N. Shinn et al., “Reflexion: Language agents with verbal reinforcement learning,” arXiv:2303.11366v4, 2023.','https://arxiv.org/abs/2303.11366v4'),
    ('C. Latimer et al., “Hindsight is 20/20: Building agent memory that retains, recalls, and reflects,” arXiv:2512.12818v1, 2025.','https://arxiv.org/abs/2512.12818v1'),
    ('T. Barnes, “Observational memory: 95% on LongMemEval,” Mastra Research, Feb. 9, 2026. Online project report; accessed Oct. 24, 2026.','https://mastra.ai/research/observational-memory'),
    ('D. Yadav, “The token-efficient memory algorithm now has temporal reasoning,” Mem0, May 14, 2026. Online project report; accessed Oct. 24, 2026.','https://mem0.ai/blog/the-token-efficient-memory-algorithm-now-has-temporal-reasoning'),
    ('D. Wu et al., “LongMemEval-V2: Evaluating long-term agent memory toward experienced colleagues,” arXiv:2605.12493v1, 2026.','https://arxiv.org/abs/2605.12493v1'),
    ('“Memory-space validation,” internal Iron Web project report, Oct. 2026, unpublished. Companion validation records: “Nemotron Lightning / Ultra: repaired and tested,” Oct. 2026.',''),
]


def reference(n,text,url):
    para=p(f"[{n}] {text}",size=12,align=WD_ALIGN_PARAGRAPH.LEFT,after=6)
    para.paragraph_format.line_spacing=1.0
    para.paragraph_format.left_indent=Inches(.32);para.paragraph_format.first_line_indent=Inches(-.32)
    para.paragraph_format.keep_together=True
    if url:
        # Hyperlink the complete bibliographic entry rather than printing long URLs.
        rel=doc.part.relate_to(url,"http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",is_external=True)
        hyperlink=OxmlElement("w:hyperlink");hyperlink.set(qn("r:id"),rel)
        for run in list(para._p.findall(qn("w:r"))):para._p.remove(run);hyperlink.append(run)
        para._p.append(hyperlink)

for i,(text,url) in enumerate(refs[:8],1):reference(i,text,url)

# PAGE 9
newpage()
h("References continued")
for i,(text,url) in enumerate(refs[8:],9):reference(i,text,url)

props=doc.core_properties
props.title="Iron Web Temporal Memory with Source Evidence"
props.subject="Nine-page enlarged research-paper editing draft"
props.author="Author Name"
props.keywords="agent memory, temporal knowledge graphs, evidence provenance"
props.comments="15-point body text is 1.5 times a 10-point baseline. Expanded one-column editing layout; final venue formatting and public experiments pending."
doc.save(DOCX)

words=len(ABSTRACT.split())
manifest={"docx":str(DOCX),"target_pages":9,"layout":"single-column enlarged editing draft",
          "body_font_points":15,"body_scale_vs_10pt":1.5,"reference_font_points":12,
          "abstract_word_count":words,"keywords":["agent memory","temporal knowledge graphs","evidence provenance"],
          "reference_count":len(refs),"author_fields":"placeholders","public_iron_web_results":"not measured"}
(OUT/"manuscript-manifest.json").write_text(json.dumps(manifest,indent=2),encoding="utf-8")
print(json.dumps(manifest,indent=2))
