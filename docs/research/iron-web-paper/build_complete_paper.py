"""Build a new, evidence-bounded manuscript without importing application code."""
from pathlib import Path
import csv
import hashlib
import json
import re
import shutil
import textwrap
from datetime import datetime, timezone

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image, ImageDraw, ImageFont
from complete_paper_content import TITLE, ABSTRACT, KEYWORDS, CONTENT, REFS

BASE = Path(__file__).resolve().parent
ROOT = BASE.parents[2]
OUT = BASE / 'manuscript' / 'reconstructed'
OUT.mkdir(parents=True, exist_ok=True)
ORIGINAL = BASE / 'manuscript' / 'Iron_Web_Research_Paper_9_Page_Editing_Draft.docx'
ORIGINAL_SHA = 'a47c2819f40da69d200cfe027763785ed0487cb70972f8f252517de30ea0c33f'
assert hashlib.sha256(ORIGINAL.read_bytes()).hexdigest() == ORIGINAL_SHA
RUNS = json.loads((BASE / 'audit-artifacts/retained_run_summary.json').read_text())

# Refresh identity checks without application imports, store access, or inference.
identity = []
with (BASE / 'audit-artifacts/artifact_hashes.csv').open(encoding='utf-8-sig', newline='') as handle:
    rows = list(csv.DictReader(handle))
for row in rows:
    target = ROOT / row['path']
    actual = hashlib.sha256(target.read_bytes()).hexdigest()
    identity.append({'path': row['path'], 'sha256': actual, 'matches_audit': actual == row['sha256']})
assert all(x['matches_audit'] for x in identity), 'Audit snapshot changed; reassess claims.'

order = []
def resolve(value):
    def cite(match):
        key = match.group(1)
        assert key in REFS
        if key not in order:
            order.append(key)
        return f'[{order.index(key)+1}]'
    return re.sub(r'\[cite:([^\]]+)\]', cite, value)

paragraphs = []
resolved = []
layout_content = [item for item in CONTENT if item[0] != 'figure']
figure_index = next(i for i,item in enumerate(layout_content) if item[0]=='p' and item[1]=='A3')
layout_content.insert(figure_index, ('figure','F1'))
for item in layout_content:
    if item[0] == 'p':
        text = resolve(item[2])
        paragraphs.append({'id': item[1], 'text': text, 'evidence': item[3]})
        resolved.append(('p', item[1], text, item[3]))
    else:
        resolved.append(item)
assert set(order) == set(REFS)

def font(size, bold=False):
    return ImageFont.truetype('C:/Windows/Fonts/timesbd.ttf' if bold else 'C:/Windows/Fonts/times.ttf', size)

canvas = Image.new('RGB', (2100, 1010), 'white')
draw = ImageDraw.Draw(canvas)
draw.text((40, 15), 'IMPLEMENTED ENTRY PATHS — DISTINCT VALIDATION BOUNDARIES', font=font(42, True), fill='black')
lanes = [
    ('Direct engine', 'Source / claim operations\nSearch and exact recall'),
    ('Hosted workbench', 'Optional extraction and answers\nSemantic + exact-field routing'),
    ('Shared context / work', 'Context cards, proofs, dependency checks\nSQLite receipts / replay on this path'),
    ('Wayfinder application', 'Separate approval and proof boundary\nApplication orchestration'),
]
for i, (name, desc) in enumerate(lanes):
    top = 100 + i * 205
    draw.rounded_rectangle((40, top, 1340, top+165), radius=12, outline='black', width=3)
    draw.text((64, top+17), name, font=font(42, True), fill='black')
    draw.multiline_text((64, top+72), desc, font=font(38), spacing=5, fill='black')
    y = top+82
    draw.line((1340, y, 1445, y), fill='black', width=4)
    draw.polygon([(1445,y),(1427,y-11),(1427,y+11)], fill='black')
draw.rounded_rectangle((1450,100,2060,880), radius=12, outline='black', width=4)
draw.multiline_text((1480,145), 'Memory engine\nand claim lifecycle', font=font(43,True), spacing=10, fill='black')
draw.multiline_text((1480,320), 'Temporal eligibility\nScoped evidence\nVector + graph retrieval\nOptional reranking', font=font(39), spacing=23, fill='black')
draw.line((1480,610,2020,610), fill='black', width=2)
draw.multiline_text((1480,650), 'FalkorDB graph\nEmbedded Qdrant\nCoordinated storage', font=font(39), spacing=21, fill='black')
draw.text((40,935), 'Arrows denote calls; evidence/results return to each caller. Live trace counts cover selected paths only.', font=font(34), fill='black')
FIG = OUT / 'architecture.png'
canvas.save(FIG)

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(.6)
section.left_margin = section.right_margin = Inches(.75)
section.header_distance = section.footer_distance = Inches(.25)
def set_font(style, size, bold=None):
    style.font.name = 'Times New Roman'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0,0,0)
    if bold is not None:
        style.font.bold = bold
    rf = style.element.get_or_add_rPr().get_or_add_rFonts()
    for key in list(rf.attrib):
        if 'Theme' in key or 'theme' in key:
            del rf.attrib[key]
    for key in ('ascii','hAnsi','cs','eastAsia'):
        rf.set(qn('w:'+key), 'Times New Roman')
for name,size,bold in [('Normal',15,False),('Title',30,True),('Subtitle',15,False),('Heading 1',19,True),('Heading 2',16,True),('Caption',12,False)]:
    set_font(doc.styles[name], size, bold)
for border in doc.styles['Title'].element.findall('.//'+qn('w:pBdr')):
    border.getparent().remove(border)
normal = doc.styles['Normal'].paragraph_format
normal.space_after = Pt(3)
normal.line_spacing = 1.0
normal.widow_control = True
for name in ['Heading 1','Heading 2']:
    pf = doc.styles[name].paragraph_format
    pf.keep_with_next = True
    pf.space_before, pf.space_after = Pt(10), Pt(6)
doc.core_properties.title = TITLE
doc.core_properties.subject = 'Architecture and implementation study with preliminary engineering validation'
doc.core_properties.author = '[Author names to be completed]'
doc.core_properties.keywords = ', '.join(KEYWORDS)
doc.core_properties.comments = 'New revision; original editing draft preserved. No new application experiments.'

def para(text, style=None):
    p = doc.add_paragraph(text, style)
    if style is None:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p
def caption(text):
    p = para(text, 'Caption')
    p.paragraph_format.keep_with_next = True
    return p
def make_table(headers, data, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    table.autofit = False
    for i,h in enumerate(headers):
        table.rows[0].cells[i].text = h
    repeat = OxmlElement('w:tblHeader')
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in data:
        for cell,value in zip(table.add_row().cells,row):
            cell.text = str(value)
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for i,cell in enumerate(row.cells):
            cell.width = Inches(widths[i])
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.space_before = Pt(4)
                for r in p.runs:
                    r.font.size = Pt(12)
                    r.bold = row == table.rows[0]
    for cell in table.rows[0].cells:
        for r in cell.paragraphs[0].runs:
            r.bold = True
    return table

p = para(TITLE, 'Title')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
for text in ['[Author names]', '[Affiliations, city, country and email addresses]', '[Funding statement and acknowledgments: authors to complete]']:
    p = para(text)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        run.font.size = Pt(12)
para('Abstract', 'Heading 1')
# Preserve the requested approximately 200-word abstract, with exactly 200 tokens by whitespace.
abstract = ABSTRACT.replace('represents explicit claim lifecycles', 'represents claim lifecycles').replace('It supports selected engineering behaviors', 'It supports engineering behaviors')
assert len(abstract.split()) == 200
para(abstract)
p = para('Keywords: ' + ', '.join(KEYWORDS))
for r in p.runs:
    r.italic = True
md = [f'# {TITLE}', '', '[Author names; affiliations; contact; funding: authors to complete]', '', '## Abstract', '', abstract, '', '**Keywords:** '+', '.join(KEYWORDS), '']
table_audit = []
for item in resolved:
    kind = item[0]
    if kind in ('h1','h2'):
        para(item[1], 'Heading 1' if kind=='h1' else 'Heading 2')
        md += [('## ' if kind=='h1' else '### ')+item[1], '']
    elif kind == 'p':
        para(item[2])
        md += [item[2], '']
    elif kind == 'figure':
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(FIG), width=Inches(7))
        figure_caption = 'Figure 1. Implemented call paths and shared memory core. The SQLite journal belongs to the shared-service path. This implementation map does not imply that every path was exercised in the retained live runs.'
        p = para(figure_caption, 'Caption')
        md += ['![Implemented architecture](architecture.png)', '', figure_caption, '']
    elif kind == 'table' and item[1] == 'T1':
        headers = ['Run','Pass','Fail','Unsupported','Disposition']
        descriptions = ['Mixed direct/hosted checks','Repeated development scenarios','Partial JSON; copied CSV excluded','Repeated final scenarios']
        data = []
        for i,run in enumerate(RUNS):
            o = run['json_outcomes']
            data.append([f'R{i+1}', o.get('Pass',0), o.get('Fail',0), o.get('Unsupported',0),descriptions[i]])
        cap = 'Table 1. Retained synthetic engineering checks, using reconciled JSON records [3]. Counts are checks, not answer accuracy or independent trials. Run identifiers appear in Appendix A.'
        caption(cap)
        make_table(headers,data,[.65,.55,.55,1.35,3.9])
        table_audit.append({'table':'T1','headers':headers,'rows':data,'source':'audit-artifacts/retained_run_summary.json','interpretation':'Engineering checks; not benchmark accuracy.'})
        md += [cap,'','| '+' | '.join(headers)+' |','|'+'---|'*len(headers)]
        md += ['| '+' | '.join(map(str,row))+' |' for row in data] + ['']
    elif kind == 'table' and item[1] == 'T2':
        headers = ['Model','Extraction (ms), n=13','Generation (ms), n=7','Adapter query (ms), n=7']
        models = ['nvidia/nemotron-3.5-lightning-30b-a3b','nvidia/nemotron-3-ultra-550b-a55b']
        data=[]
        for label,model in zip(['Lightning','Ultra'],models):
            metrics = RUNS[-1]['observed_timings'][model]
            data.append([label]+[f"{metrics[key]['median']:.3f}" for key in ('extraction_inference_ms','generation_inference_ms','query_adapter_total_ms')])
        cap = 'Table 2. Medians of recorded R4 hosted-path timings [3]. One replayed extraction trace per model is excluded; recorded repair attempts remain included. These samples have no matched architecture baseline.'
        caption(cap)
        make_table(headers,data,[1,2,2,2])
        p = para('Model identifiers: '+models[0]+' (Lightning); '+models[1]+' (Ultra).', 'Caption')
        table_audit.append({'table':'T2','headers':headers,'rows':data,'model_ids':models,'source':'audit-artifacts/retained_run_summary.json','interpretation':'Descriptive medians in milliseconds, not controlled comparative effects.'})
        md += [cap,'','| '+' | '.join(headers)+' |','|'+'---|'*len(headers)]
        md += ['| '+' | '.join(row)+' |' for row in data]+['','Model identifiers: '+'; '.join(models),'']

para('References','Heading 1')
md += ['## References','']
for i,key in enumerate(order,1):
    ref = REFS[key]
    entry = f'[{i}] '+ref['text']
    if ref['url']:
        entry += ' '+ref['url']
    p = para(entry)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Inches(.25)
    p.paragraph_format.first_line_indent = Inches(-.25)
    for r in p.runs:
        r.font.size=Pt(12)
    md += [entry,'']
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('Iron Web | Architecture and implementation study | ')
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'),'PAGE')
footer._p.append(field)
for r in footer.runs:
    r.font.size=Pt(10)
target=OUT/'Iron_Web_Complete_Research_Paper.docx'
doc.save(target)
(OUT/'Iron_Web_Complete_Research_Paper.md').write_text('\n'.join(md),encoding='utf-8')
(OUT/'paper_claim_map.json').write_text(json.dumps({'paragraphs':paragraphs,'tables':table_audit},indent=2,ensure_ascii=False),encoding='utf-8')
(OUT/'reference_verification.json').write_text(json.dumps([{**REFS[key],'key':key,'number':i+1} for i,key in enumerate(order)],indent=2,ensure_ascii=False),encoding='utf-8')

audit = '''# Manuscript claim-to-evidence audit

Scope: complete manuscript revision, 24 October 2026. This report is separate from the paper. Paragraph-level text and evidence keys are retained in `paper_claim_map.json`; bibliography verification is in `reference_verification.json`. Evidence paths are repository-relative unless they name an audit file in the existing research folder. Source inspection is not a new execution.

| Claim / manuscript location | Supporting evidence | Status and limits |
|---|---|---|
| Temporal claim lifecycle and known-time eligibility (3; 4) | `memory/lifecycle/models.py`, `memory/lifecycle/service.py`; implementation audit | Implemented. Eligibility is represented policy, not verified truth; correction is not ordinary historical supersession. |
| Four distinct entry paths (Fig. 1) | Engine, workbench, shared service, Wayfinder; implementation audit | Diagram is an implementation map. Hosted live runs do not exercise the shared context/work service. |
| Scoped context proofs, post-generation check, dependent-work traversal (4; 7.2) | `memory/context/access.py`, `memory/shared/client.py`, lifecycle invalidation/handoff; targeted test source | Implemented with precise regression assertions. End-to-end stale acceptance Not measured; no guarantee for external side effects. |
| Vector reconciliation and confounded ablations (7.2; 8) | Vector/hybrid retrievers and engine mode dispatch | No raw dense, fully graph-free, or fully temporal-free baseline is established by those labels. |
| Defaults and hardware (5) | `config.yaml`, config/engine/shared constructors; audited environment metadata | CPU defaults and configured models are inspected. Historical resolved manifests absent; GPU is user-reported. |
| R1: 11 Pass / 10 Fail / 3 Unsupported (Table 1) | `docs/live-validation/live-test-evidence.json`, reconciled summary | 24 synthetic checks. Mixed paths, correlated outcomes; not benchmark accuracy. |
| R2: 13 Pass / 5 Fail (Table 1) | `docs/nemotron-validation-before-schema/nemotron-evidence.json`, summary | 18 development checks, including correction, extraction, and citation failures. |
| R3: 2 Pass / 6 Fail (Table 1) | `docs/nemotron-validation-schema-rejected/nemotron-evidence.json` | Partial eight-check JSON retained. Its copied 18-row CSV is identical to R2; copied summary excluded as R3 evidence. |
| R4: 18 Pass / 0 Fail (Table 1) | `docs/nemotron-validation/nemotron-evidence.json`, summary | Repeated repaired scenarios across two models; not held-out generalization. |
| Timings (Table 2) | R4 event traces; `retained_run_summary.json` | Exact medians recomputed previously; n=13 extraction, n=7 generation/query per model, one replay trace excluded each. Repair included; no cost or tail-latency claim. |
| External methods (2) | Versioned primary papers and standards listed in reference verification | Conceptual comparison only. No heterogeneous leaderboard; no claim that an undocumented feature is absent. |
| Storage coordination and recovery (4; 9) | `memory/storage.py`, journal/snapshot code, audited historical recovery records | Compensation is not a transaction spanning all stores; recovery metadata constrains broad durability claims. |
| Source identity and reproducibility (Appendix A) | Base commit plus 202-file hash manifest and 833-symbol index | All 202 hashes rechecked unchanged while generating this revision. Historical runtimes remain incompletely reproducible. |

## Claims deliberately removed or weakened

- The illustrative 98% accuracy and claims of superiority have no experiment behind them and are omitted.
- Unit-test pass counts are not memory accuracy. Historical suite totals overlap and are not summed or promoted to a headline result.
- Perfect historical recall, universal contradiction detection, complete stale-work protection, calibrated truth probabilities, atomic multi-store writes, complete erasure, adversarial isolation, and lower token cost are not asserted.
- Graph expansion and fuzzy scoring are implemented, but their incremental utility is Not measured.
- Copied R3 summaries and unexecuted planned benchmark cases are not counted as completed observations.
- Current source is not treated as the exact source revision for every historical trace.

## Checks performed for this revision

The original DOCX hash is unchanged. Citation numbering follows first appearance and every bibliography entry is cited. Table 1 is populated from JSON outcomes, retaining failures; Table 2 is populated from the audited timing summary. No application code imports, inference APIs, benchmarks, tests, or memory database operations were needed. Document structure, embedded diagram, table cells, text extraction, and rendered pages are checked separately in the generation manifest.
'''
(OUT/'paper_claim_evidence_audit.md').write_text(audit,encoding='utf-8')
notes = f'''# Paper generation notes

New manuscript: `Iron_Web_Complete_Research_Paper.docx`. The Markdown sibling is an editable text source; the Word document contains the figure and formatted tables. The accompanying PDF is produced with installed Microsoft Word and inspected after rendering.

## Scope and format

This is the requested complete paper, with all ten numbered sections, an approximately 200-word abstract (200 whitespace-delimited words), three keywords, architecture diagram, two evidence-backed tables, numbered references, and a reproducibility appendix. It is an architecture and implementation study with preliminary engineering validation. Body text is 15-point Times New Roman, retaining the earlier request for 1.5-times editing text. References and tables are 12 point. The latest complete-manuscript structure takes precedence over forcing the expanded editing copy into the earlier nine-page layout. This is an editing manuscript, not a claim of conformance to a specific IEEE conference template.

Original backup remains untouched at `../Iron_Web_Research_Paper_9_Page_Editing_Draft.docx`; SHA-256 `{ORIGINAL_SHA}`. No application implementation was changed for this revision. Existing unrelated working-tree changes were preserved.

## Missing evidence and provisional material

- No eligible controlled head-to-head public benchmark, held-out conversational experiment, matched graph ablation, or causal reliability effect size is available.
- Full historical resolved configurations, seeds/repetition policies, reader token budgets, and provider/hardware manifests are missing. The declared GPU is user-reported.
- Named `regression.xml`, `startup-memory-tests.xml`, `phase02-tests.xml`, `real-embedding-smoke.json`, and `memory-spaces-live-evidence.json` are absent from the accessible project. `.local/tests-relocation-20261024` was inaccessible in the audit.
- Shared-service context/work execution is implemented but is not exercised by retained hosted runs. End-to-end stale acceptance and false-block rates remain Not measured.
- Storage scaling, robust tail latency, full token/cost accounting, crash coverage, and adversarial behavior remain unmeasured.
- R3 includes an incompatible copied CSV and stale report. The negative partial JSON is retained; the copied summary is excluded from quantitative attribution.
- Proposed comparisons in Methodology and Future Work are explicitly unexecuted. Author, affiliation, email, funding and acknowledgment fields are editable metadata requiring the authors' completion. No invented results or unfinished result placeholders occur.
- A sanitized public artifact package has not been released. The internal report reference points to real local artifacts, not a peer-reviewed public publication.

## Citation verification

All twelve entries were checked against primary publication/version records, official standards, author-institution metadata, or actual local artifacts. No critical citation identity check remains unresolved. Verification scope is recorded individually in `reference_verification.json`. The author-hosted full text of the 1986 temporal-database paper timed out; only the verified institutional metadata/abstract supports the general temporal-dimension precedent. Hindsight is cited only for its documented conceptual distinction, not an unverified numerical result. Zep's official record and supplied v1 PDF were used when its versioned web record did not load. Preprint versions retain their actual version years rather than inventing conference pages.

## Rebuilding and rendering

Run `build_complete_paper.py` with the bundled dependency Python; it imports only document/data libraries and the pure manuscript-content module. It verifies all 202 audited file hashes before building. It must not import the earlier `build_manuscript.py`, which writes the original draft at import time. The document skill's standard renderer requires LibreOffice, unavailable on this host; installed Word provides local PDF export via `render_with_word.ps1`. PDF pages are rasterized with PDFium for visual inspection. Rendering uses no inference service and opens no memory store.

The paragraph-level mapping and table source records remain in `paper_claim_map.json`. Final page count, checks, hashes, and render inspection are recorded in `generation_manifest.json`.
'''
(OUT/'paper_generation_notes.md').write_text(notes,encoding='utf-8')
manifest = {'generated_at_utc':datetime.now(timezone.utc).isoformat(),'original_sha256':ORIGINAL_SHA,'original_unchanged':True,'audited_files_rechecked':len(identity),'snapshot_unchanged':all(x['matches_audit'] for x in identity),'abstract_words':len(abstract.split()),'body_words':sum(len(p['text'].split()) for p in paragraphs),'references':len(order),'main_sections':10,'figures':1,'tables':2,'new_experiments':False,'application_modified':False,'databases_opened':False,'pdf_render':'pending','visual_inspection':'pending'}
(OUT/'generation_manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
print(json.dumps({'output':str(target),**manifest},indent=2))
